"""KPI Analysis · Report Export: one KPI over an area and a period.

The report is judged by the KPI Analysis engine with the user's threshold and
the KPI's own direction, and by each cell's own hours for sudden spikes; the
charts count only the affected cells; the Excel workbook is the Draw Data
export itself, of the affected cells; the trend and the Top Sites chart are
hourly; there is no City analysis; the ticket table has the City, the SLA
Target Time and the KPI value; the PowerPoint is real slides with native
charts, the Huawei logo and none of the old slogans; the files are named after
what is in them and land on the Desktop (a test folder here).
"""

import io
import sys
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "app"
if str(APP) not in sys.path:
    sys.path.insert(0, str(APP))

PRB = "HW_DL PRB Avg Utilization(%)"
INTER = "L.UL.Interference.Avg(dBm)"
AVA = "LTE_Availability(%)@AB"
FLOW = "VS.RscGroup.FlowCtrol.DL.DropNum"
SLOGANS = ("Building a Fully Connected", "Intelligent World", "Iraq Network Optimization",
           "Italian Bridge", "Network Insights for a Better Tomorrow")
TICKET_HEAD = ["Ticket ID", "Site ID", "City", "SLA Target Time", "Create Time",
               "Problem Time", "Is CMC", "KPI Value (dBm)"]


def _export() -> bytes:
    """BAS0001: an FDD cell and an interfered TDD cell; BAS0002: two TDD cells,
    one of them just above -105 dBm, the other clean."""
    head = f"Time,eNodeB Name,Cell FDD TDD Indication,Cell Name,LocalCell Id,{PRB},{INTER},{AVA}\n"
    rows = [head]
    for day in ("2026-09-13", "2026-09-14"):
        for hh in range(24):
            t = f"{day} {hh:02d}:00"
            rows.append(f"{t},Alpha_BAS0001,CELL_FDD,L_Alpha_BAS0001-1,1,90,-118,100\n")
            rows.append(f"{t},Alpha_BAS0001,CELL_TDD,T_Alpha_BAS0001-2,2,30,-99,100\n")
            rows.append(f"{t},Beta_BAS0002,CELL_TDD,T_Beta_BAS0002-1,1,20,-104,97\n")
            rows.append(f"{t},Beta_BAS0002,CELL_TDD,T_Beta_BAS0002-2,2,20,-112,100\n")
    return "".join(rows).encode()


def _flow_export() -> bytes:
    """48 hours of 3G flow-control drops, the user's case: two NodeBs that run
    at 20-40 an hour and jump once (to 100,000 and to 600,000), one that is
    always around 65,000, one whose jump stays small (800), one at zero."""
    import zipfile
    lines = ["\n\n\nSHAMS-3G\nSave Time :2026-09-11 10:51:53\n\n",
             f"Time,RNC,NODEBNAME,NodeB ID,Integrity,{FLOW},VS.IPPM.Rtt.Means(ms),"
             "3G_Availability@AB\n"]
    for h in range(48):
        t = f"2026-09-{8 + h // 24:02d} {h % 24:02d}:00"
        base = (20, 30, 40)[h % 3]
        for name, v in (("Spiky_BAS0101", 100_000 if h == 30 else base),
                        ("Huge_BAS0102", 600_000 if h == 31 else base + 5),
                        ("Chronic_BAS0103", 60_000 + 5_000 * (h % 3)),
                        ("Small_BAS0104", 800 if h == 32 else base),
                        ("Calm_BAS0105", 0)):
            lines.append(f"{t},RBASH01,{name},1,100%,{v},3,100\n")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("export(Subreport 1).csv", "".join(lines))
    return buf.getvalue()


def _store(put_resource):
    from test_complaint_analysis import _target_bytes
    from rfopt.complaints.target_store import save_target
    put_resource("kpi", "R5 4G Monitoring Hourly KPI.csv", _export(), "4G KPI")
    save_target(_target_bytes(), "Target 13-Sep.xlsx")


def _rows(values: dict, hours: int = 6) -> pd.DataFrame:
    """object -> (site, duplex, value per hour) as the KPI export has them."""
    t0 = pd.Timestamp("2026-09-13 00:00")
    out = []
    for obj, (site, duplex, vals) in values.items():
        for h in range(hours):
            out.append({"datetime": t0 + pd.Timedelta(hours=h), "object": obj, "site_id": site,
                        "duplex": duplex, "parent": f"eNB_{site}",
                        INTER: vals[h] if isinstance(vals, list) else vals, "tech": "4G"})
    return pd.DataFrame(out)


def _report(**kw):
    from rfopt.reports.kpi_report import KpiReport
    cells = pd.DataFrame({
        "object_id": ["A-1", "A-2", "B-1", "B-2", "C-1"],
        "site_id": ["BAS0001", "BAS0001", "BAS0002", "BAS0002", "NAS0001"],
        "cell_name": ["A-1", "A-2", "B-1", "B-2", "C-1"],
        "governorate": ["Basrah", "Basrah", "Basrah", "Basrah", "Dhi Qar"],
        "sup_district": ["Markaz Al-Zubair", "Markaz Al-Zubair", "Safwan", "Safwan",
                         "Markaz Al-Nasiriya"],
        "value": [-99.0, -110.0, -104.0, -103.0, -101.0],
        "peak": [-98.0, -109.0, -103.0, -102.0, -100.0],
        "peak_time": pd.to_datetime(["2026-09-13 10:00"] * 5),
        "worst": [-98.0, -109.0, -103.0, -102.0, -95.0],
        "beyond": [True, False, True, True, True],
        "spike": [False, False, False, False, True],
        "spike_value": [np.nan] * 4 + [-95.0],
        "spike_time": pd.to_datetime([None] * 4 + ["2026-09-13 03:00"]),
    })
    cells["affected"] = cells["beyond"] | cells["spike"]
    rows = _rows({"A-1": ("BAS0001", "CELL_TDD", -99.0), "A-2": ("BAS0001", "CELL_TDD", -110.0),
                  "B-1": ("BAS0002", "CELL_TDD", -104.0), "B-2": ("BAS0002", "CELL_FDD", -103.0),
                  "C-1": ("NAS0001", "CELL_TDD", [-101.0, -101.0, -101.0, -95.0, -101.0, -101.0])})
    tickets = pd.DataFrame({"Ticket ID": ["T1", "T2", "T3"],
                            "Site ID": ["BAS0002", "BAS0002", "BAS0001"],
                            "City": ["Basrah", "Basrah", "Zubair"],
                            "SLA Target Time": pd.to_datetime(["2026-09-15 10:30"] * 3),
                            "Create Time": pd.to_datetime(["2026-09-13 10:30"] * 3),
                            "Problem Time": pd.to_datetime(["2026-09-13 01:00", "2026-09-13 02:00",
                                                            "2026-09-13 03:00"]),
                            "Is CMC": ["Yes", "No", "Yes"],
                            "KPI Value": [-103.5, -103.5, -104.5]})
    args = dict(tech="4G", kpi="4G TDD Interference", column=INTER, unit=" dBm",
                low_is_bad=False, per_day=False, threshold=-105.0, governorate="Basrah",
                sup_districts=["Markaz Al-Zubair", "Safwan", "Um Qasr"], site="",
                start=date(2026, 9, 13), end=date(2026, 9, 14), cells=cells, rows=rows,
                trend=rows.groupby("datetime")[INTER].mean(), tickets=tickets)
    args.update(kw)
    return KpiReport(**args)


def _pptx(r):
    from pptx import Presentation
    from _kpi_report import ASSETS
    from rfopt.reports.kpi_report import Assets, build_pptx
    return Presentation(io.BytesIO(build_pptx(r, Assets(ASSETS / "report_cover.jpg",
                                                        ASSETS / "huawei_logo_white.png"))))


def _texts(slide) -> str:
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def _ws(tech: str):
    import _kpi_workspace as W
    import _resources as R
    files = W.combine(R.kpi_groups())
    usable = [(b, i) for b, i in files if i.kind == tech]
    index = pd.concat([W.index_of(b) for b, i in usable], ignore_index=True)
    return W.Workspace(files, usable, tech, "NodeB", index, "Network", None, [], set(), [],
                       None, None)


# --------------------------------------------------------------------------- #
# the judgement: the threshold, and sudden spikes against a cell's own hours
# --------------------------------------------------------------------------- #
def test_a_sudden_spike_is_found_against_the_cells_own_normal_hours():
    """Flow control: 20-40 an hour that jumps to 100,000 is a spike, and so is
    one that jumps above 500,000; an object always around 65,000 is not (the
    threshold judges it), nor a jump that stays small, nor a daily busy hour."""
    from rfopt.kpi.anomaly import sudden_spikes
    t = pd.date_range("2026-09-08", periods=48, freq="h")
    base = np.array([20.0, 30.0, 40.0] * 16)

    def frame(name, vals):
        return pd.DataFrame({"object": name, "datetime": t, FLOW: vals})

    spiky, huge, small = base.copy(), base.copy() + 5, base.copy()
    spiky[30], huge[31], small[32] = 100_000, 600_000, 800
    busy = np.where(t.hour == 20, 45_000.0, 40_000.0)             # a daily busy hour
    df = pd.concat([frame("Spiky", spiky), frame("Huge", huge), frame("Small", small),
                    frame("Chronic", 60_000 + 5_000 * (np.arange(48) % 3)),
                    frame("Busy", busy), frame("Calm", np.zeros(48))])
    s = sudden_spikes(df, FLOW, low_is_bad=False, level=50_000).set_index("object")
    assert sorted(s.index) == ["Huge", "Spiky"]
    assert s.loc["Spiky", "spike_value"] == 100_000 and s.loc["Spiky", "normal"] == 30
    assert s.loc["Spiky", "spike_time"] == t[30] and s.loc["Huge", "spike_value"] == 600_000
    # the other way: availability that collapses from 100 is a sudden drop
    ava = pd.DataFrame({"object": "Cell", "datetime": t,
                        AVA: np.where(np.arange(48) == 10, 0.0, 100.0)})
    d = sudden_spikes(ava, AVA, low_is_bad=True, level=99.9)
    assert list(d["object"]) == ["Cell"] and d["spike_value"].iloc[0] == 0.0
    # too few hours to know a normal: nothing is called a spike
    short = pd.DataFrame({"object": "Short", "datetime": t[:3], FLOW: [30, 30, 100_000]})
    assert sudden_spikes(short, FLOW, low_is_bad=False, level=50_000).empty


def test_the_report_finds_a_flow_control_spike_the_threshold_misses(put_resource):
    """At the configured 500,000 per 24 h, the NodeB that jumped from 20-40 to
    100,000 is affected through its spike; the one always near 65,000 through
    the threshold; the small jump and the calm NodeB are not."""
    import _kpi_report as X
    put_resource("kpi", "SHAMS-3G.zip", _flow_export(), "3G KPI")
    ws = _ws("3G")
    ch = next(c for c in X.kpi_choices(ws) if c.column == FLOW)
    assert ch.judged.how == "day_sum" and ch.judged.rule.critical == 500_000
    r = X._build(("flow",), ws, ch, "", (), "", (date(2026, 9, 8), date(2026, 9, 9)),
                 500_000.0)
    c = r.cells.set_index("object_id")
    assert set(r.affected["object_id"]) == {"Spiky_BAS0101", "Huge_BAS0102", "Chronic_BAS0103"}
    assert not c.loc["Spiky_BAS0101", "beyond"] and c.loc["Spiky_BAS0101", "spike"]
    assert c.loc["Chronic_BAS0103", "beyond"] and not c.loc["Chronic_BAS0103", "spike"]
    assert c.loc["Huge_BAS0102", "beyond"] and c.loc["Huge_BAS0102", "spike"]
    assert list(r.top_sites()["site_id"]) == ["BAS0102", "BAS0101", "BAS0103"]
    assert "or showed a sudden spike" in r.description
    assert r.per_day and len(r.trend) == 48                 # hourly, never folded into days


# --------------------------------------------------------------------------- #
# the report's numbers
# --------------------------------------------------------------------------- #
def test_the_charts_count_only_the_affected_cells():
    r = _report()
    assert (r.total_sites, r.affected_sites, r.total_cells, r.affected_cells) == (3, 3, 5, 4)
    assert r.spike_cells == 1
    # every picked Sup District is listed, a clean one with 0
    assert r.affected_by("sup_district", ("Markaz Al-Zubair", "Safwan", "Um Qasr")).to_dict() == {
        "Safwan": 2, "Markaz Al-Nasiriya": 1, "Markaz Al-Zubair": 1, "Um Qasr": 0}
    assert r.affected_by("governorate").to_dict() == {"Basrah": 3, "Dhi Qar": 1}
    # interference: the higher the worse — by the worst hour; A-2 is never in it
    top = r.top_sites()
    assert list(top["site_id"]) == ["NAS0001", "BAS0001", "BAS0002"]
    assert list(top["worst"]) == [-95.0, -98.0, -102.0]
    assert r.top_ticket_sites().to_dict() == {"BAS0002": 2, "BAS0001": 1}
    assert r.total_tickets == 3


def test_the_top_sites_lines_are_hourly_as_draw_data_draws_a_site():
    """A site's hour is its cells averaged for a level, the way Draw Data draws a
    site (`rfopt.kpi.trends.series_for`)."""
    from rfopt.kpi.trends import series_for
    r = _report()
    lines = r.site_lines(10)
    assert list(lines) == ["NAS0001", "BAS0001", "BAS0002"]
    bas1 = lines["BAS0001"]
    assert len(bas1) == 6 and bas1.index.freq is None and (bas1.diff().dropna() == 0).all()
    assert bas1.iloc[0] == pytest.approx((-99.0 - 110.0) / 2)
    pd.testing.assert_series_equal(bas1, series_for(r.rows, INTER, level="Site", obj="BAS0001"))
    assert list(lines["NAS0001"].index) == list(pd.date_range("2026-09-13", periods=6, freq="h"))


def test_a_low_is_bad_kpi_is_judged_the_other_way():
    r = _report(kpi="4G Availability", unit="%", low_is_bad=True, threshold=99.0)
    r.cells["value"] = [100.0, 98.0, 97.0, 99.5, 100.0]
    r.cells["worst"] = [100.0, 98.0, 97.0, 99.5, 100.0]
    r.cells["beyond"] = r.cells["value"] < 99.0
    r.cells["spike"] = False
    r.cells["affected"] = r.cells["beyond"]
    top = r.top_sites()
    assert list(top["site_id"]) == ["BAS0002", "BAS0001"] and list(top["worst"]) == [97.0, 98.0]
    assert r.threshold_text == "< 99%" and "fell below" in r.description
    assert "sudden" not in r.description                    # no spike: the old sentence


def test_the_cover_summary_has_no_threshold_and_the_description_stays_short():
    r = _report()
    rows = dict(r.summary_rows())
    assert list(rows) == ["Governorate", "Sup District(s)", "Site", "Technology", "KPI",
                          "Period", "Total Sites", "Affected Sites", "Total Cells",
                          "Affected Cells", "Total Tickets"]
    assert rows["Affected Cells"] == "4 (80.0%)" and rows["Total Tickets"] == "3"
    assert "Threshold" not in rows and "-105" not in " ".join(rows.values())
    assert r.description.startswith("4G TDD Interference exceeded the configured threshold "
                                    "or showed a sudden spike on 4 of 5")
    assert len(r.description) < 260


def test_the_file_names_say_what_is_in_them():
    from rfopt.reports.kpi_report import pptx_name, xlsx_name
    r = _report(sup_districts=["Zubair"])
    assert pptx_name(r, date(2026, 9, 18)) == \
        "KPI_Report_4G_TDD_Interference_Basrah_Zubair_2026-09-18.pptx"
    assert xlsx_name(r, date(2026, 9, 18)) == \
        "KPI_Report_Data_4G_TDD_Interference_Basrah_Zubair_2026-09-18.xlsx"
    odd = _report(governorate="Basrah", sup_districts=['Al:Siba/"x"', "B", "C", "D"])
    name = pptx_name(odd, date(2026, 9, 18))
    assert not set('\\/:*?"<>|') & set(name) and "plus2" in name


# --------------------------------------------------------------------------- #
# the Excel workbook: the Draw Data export, of the affected cells
# --------------------------------------------------------------------------- #
def _book(data_or_path):
    from openpyxl import load_workbook
    src = io.BytesIO(data_or_path) if isinstance(data_or_path, bytes) else data_or_path
    return load_workbook(src)


def _layout(ws) -> dict:
    return {"merged": sorted(str(m) for m in ws.merged_cells.ranges),
            "freeze": ws.freeze_panes, "filter": ws.auto_filter.ref,
            "widths": {k: d.width for k, d in ws.column_dimensions.items()},
            "rules": sorted((str(cf.sqref), r.type, str(r.operator), str(r.formula),
                             bool(r.stopIfTrue)) for cf in ws.conditional_formatting
                            for r in cf.rules),
            "fonts": [(c.font.name, c.font.b, c.fill.fgColor.rgb, c.number_format)
                      for row in ws.iter_rows(min_row=1, max_row=3) for c in row]}


def test_the_workbook_is_the_draw_data_export_of_the_affected_cells(tmp_path):
    from rfopt.reports.kpi_pivot import write_pivot_workbook
    from rfopt.reports.kpi_report import build_xlsx
    r = _report()
    mine = _book(build_xlsx(r))
    path, sheets, _ = write_pivot_workbook(r.excel_rows(), [INTER], tmp_path / "draw.xlsx")
    ref = _book(path)
    assert mine.sheetnames == ref.sheetnames == ["UL Inter"] == [r.sheet]
    a, b = mine["UL Inter"], ref["UL Inter"]
    assert list(a.iter_rows(values_only=True)) == list(b.iter_rows(values_only=True))
    assert _layout(a) == _layout(b)
    # the time runs across: a row per affected cell, a column per hour
    head = [c.value for c in a[2]]
    assert head[:2] == ["Row Labels", "Cell FDD TDD Indication"]
    assert head[2:] == [f"09/13 {h:02d}:00 AM".replace("00:00 AM", "12:00 AM")
                        for h in range(6)]
    assert a["A1"].value.startswith(f"Average of {INTER} - 2026-09-13 - Hourly")
    labels = [a.cell(i, 1).value for i in range(3, a.max_row + 1)]
    assert labels == ["A-1", "B-1", "B-2", "C-1"]            # A-2 is not affected


def test_the_workbook_matches_draw_datas_own_export(put_resource, tmp_path, monkeypatch):
    """The same file through Draw Data's Excel export (`_kpi_workspace.export`)
    and through the report: same sheet, same banner and header, and each of the
    report's rows is Draw Data's row for that cell, value for value."""
    import _kpi_report as X
    import _kpi_workspace as W
    import _shared
    from rfopt.reports.kpi_report import build_xlsx
    _store(put_resource)
    ws = _ws("4G")
    monkeypatch.setattr(_shared, "DL", tmp_path)            # never the user's Downloads
    note = W.export([INTER], ws.usable)
    assert "Saved to Downloads" in note
    draw = _book(next(tmp_path.glob("*.xlsx")))
    ch = next(c for c in X.kpi_choices(ws) if c.label == "4G UL interference")
    r = X._build(("cmp",), ws, ch, "", (), "", (date(2026, 9, 13), date(2026, 9, 14)),
                 -105.0)
    mine = _book(build_xlsx(r))
    assert mine.sheetnames == draw.sheetnames == ["UL Inter"]
    a, b = mine["UL Inter"], draw["UL Inter"]
    assert a["A1"].value == b["A1"].value                     # the same banner
    assert [c.value for c in a[2]] == [c.value for c in b[2]]  # the same header, 48 hours
    theirs = {row[0]: row for row in b.iter_rows(min_row=3, values_only=True)}
    ours = list(a.iter_rows(min_row=3, values_only=True))
    assert {row[0] for row in ours} == {"T_Alpha_BAS0001-2", "T_Beta_BAS0002-1"} < set(theirs)
    for row in ours:
        assert row == theirs[row[0]]
    assert _layout(a)["rules"] == [(rule[0].replace(str(len(theirs) + 2), str(len(ours) + 2)),
                                    *rule[1:]) for rule in _layout(b)["rules"]]


# --------------------------------------------------------------------------- #
# the PowerPoint
# --------------------------------------------------------------------------- #
def test_the_powerpoint_is_real_slides_with_native_charts_and_the_huawei_logo():
    prs = _pptx(_report())
    assert len(prs.slides) == 6 and round(prs.slide_width.inches, 2) == 13.33
    cover = prs.slides[0]
    pics = [s for s in cover.shapes if s.shape_type == 13]
    assert len(pics) == 2                                   # the bridge and the logo
    charts = [s.chart.chart_type for sl in prs.slides for s in sl.shapes if s.has_chart]
    assert len(charts) >= 5                                 # native, editable charts
    text = " ".join(_texts(sl) for sl in prs.slides)
    for slogan in SLOGANS:
        assert slogan not in text, slogan
    cover_text = _texts(cover)
    assert "Threshold" not in cover_text and "Analysis Report" in cover_text


def test_slide_2_is_the_hourly_network_trend():
    prs = _pptx(_report())
    trend = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
    assert [p.name for p in trend.plots[0].series] == ["Network Average", "Threshold (-105 dBm)"]
    cats = list(trend.plots[0].categories)
    assert cats == [f"13 Sep {h:02d}:00" for h in range(6)]  # every hour, as exported
    assert "(hourly)" in _texts(prs.slides[1])
    # a 24 h total is not drawn as a line on an hourly chart
    day = _pptx(_report(per_day=True, kpi="3G DL flow-control drops", unit=""))
    trend = next(s.chart for s in day.slides[1].shapes if s.has_chart)
    assert [p.name for p in trend.plots[0].series] == ["Network Average"]
    assert len(list(trend.plots[0].categories)) == 6


def test_there_is_no_city_analysis():
    prs = _pptx(_report())
    areas = prs.slides[2]
    assert len([s for s in areas.shapes if s.has_chart]) == 2
    text = _texts(areas)
    assert "by Governorate" in text and "by Sup District" in text and "City" not in text
    from _kpi_report import slide_html
    html = slide_html(_report(), 2)
    assert "by Governorate" in html and "by City" not in html


def test_slide_4_is_the_hourly_top_sites_line_chart():
    from pptx.enum.chart import XL_CHART_TYPE
    prs = _pptx(_report())
    top = prs.slides[3]
    charts = [s.chart for s in top.shapes if s.has_chart]
    assert len(charts) == 1 and charts[0].chart_type == XL_CHART_TYPE.LINE
    assert [p.name for p in charts[0].plots[0].series] == ["NAS0001", "BAS0001", "BAS0002"]
    assert list(charts[0].plots[0].categories) == [f"13 Sep {h:02d}:00" for h in range(6)]
    text = _texts(top)
    assert "Top 10 Sites by TDD Interference" in text and "Affected Cells by Site" not in text
    from _kpi_report import slide_html
    html = slide_html(_report(), 3)
    assert "Top 10 Sites by" in html and "Affected cells by Site" not in html
    assert html.count('class="rx-p"') == 1                   # the one chart, nothing else


def test_slide_6_is_received_tickets_with_city_sla_target_time_and_the_kpi():
    from rfopt.reports.kpi_report import TICKET_COLS
    prs = _pptx(_report())
    slide = prs.slides[5]
    titles = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert titles[0] == "Received Tickets"
    table = next(s for s in slide.shapes if s.has_table).table
    assert [c.text for c in table.rows[0].cells] == TICKET_HEAD
    assert TICKET_COLS[-1] == "KPI Value" and "SLA" not in TICKET_COLS
    first = [c.text for c in table.rows[1].cells]            # the latest problem first
    assert first == ["T3", "BAS0001", "Zubair", "2026-09-15 10:30", "2026-09-13 10:30",
                     "2026-09-13 03:00", "Yes", "-104.50"]
    from _kpi_report import slide_html
    html = slide_html(_report(), 5)
    assert "<b>Received Tickets</b>" in html
    assert "".join(f"<th>{h}</th>" for h in TICKET_HEAD) in html


# --------------------------------------------------------------------------- #
# the tab
# --------------------------------------------------------------------------- #
def test_the_draw_data_colours_are_the_ones_the_report_draws_with():
    from _charts import LINE_COLOURS
    from rfopt.reports import kpi_report
    assert kpi_report.LINE_COLOURS == LINE_COLOURS


def test_the_report_export_tab(put_resource):
    """The tab on the loaded exports: its filters, the live preview, the Excel
    preview (the Draw Data layout), and both exports saved (to the test's Desktop)."""
    import os
    AppTest = pytest.importorskip("streamlit.testing.v1").AppTest
    _store(put_resource)
    at = AppTest.from_file(str(APP / "views/kpi_analysis.py"), default_timeout=300)
    at.session_state["ka_view"] = "Report Export"
    at.run()
    assert not at.exception, at.exception
    labels = {w.label for w in list(at.selectbox) + list(at.multiselect)}
    assert {"Technology", "KPI", "Governorate", "Sup District", "Site"} <= labels
    assert "City" not in labels
    assert [n.label for n in at.number_input][0] == "Threshold (dBm)"
    assert at.get("date_input")[0].label == "Time Period"
    assert at.selectbox(key="rx_kpi").value == "4G TDD Interference"

    def slide():
        return " ".join(e.proto.body for e in at.get("html") if "rx-slide" in e.proto.body)

    cover = slide()
    assert 'alt="HUAWEI"' in cover and "Analysis Report" in cover
    assert "<span>Total Cells</span><span>:</span><b>3</b>" in cover
    assert "<span>Affected Cells</span><span>:</span><b>2 (66.7%)</b>" in cover
    # the Excel preview is the Draw Data sheet: a row per affected cell, the hours across
    grid = at.dataframe[0].value
    assert list(grid.columns[:2]) == ["Row Labels", "Cell FDD TDD Indication"]
    assert len(grid.columns) == 2 + 48 and list(grid["Row Labels"]) == [
        "T_Alpha_BAS0001-2", "T_Beta_BAS0002-1"]
    # the threshold moves the numbers — the preview follows
    at.number_input(key="rx_thr_4G TDD Interference").set_value(-100.0).run()
    assert "<span>Affected Cells</span><span>:</span><b>1 (33.3%)</b>" in slide()
    for k in range(2, 7):
        at.button(key="rx_next").click().run()
        assert not at.exception, at.exception
        assert f"{k} / 6" in slide()
    assert "<b>Received Tickets</b>" in slide()

    at.button(key="rx_exp_pptx").click().run()
    at.button(key="rx_exp_xlsx").click().run()
    assert not at.exception, at.exception
    desk = Path(os.environ["RFOPT_DESKTOP"])
    saved = sorted(p.name for p in desk.iterdir())
    assert len(saved) == 2 and saved[0].startswith("KPI_Report_4G_TDD_Interference_R5_")
    assert saved[1].startswith("KPI_Report_Data_4G_TDD_Interference_R5_")
    assert sum("Saved to the Desktop" in s.value for s in at.success) == 2
    book = _book(desk / saved[1])
    assert book.sheetnames == ["UL Inter"]
    assert [book["UL Inter"].cell(i, 1).value for i in range(3, book["UL Inter"].max_row + 1)] \
        == ["T_Alpha_BAS0001-2"]
