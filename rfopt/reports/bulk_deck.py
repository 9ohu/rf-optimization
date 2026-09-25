"""KPI Analysis · Draw Data — the Bulk Draw deck.

The charts Bulk Draw has drawn, six to a slide (three across, two down) with
the site over each, and the combined chart on a slide of its own at the end.
The dress is the Report Export's own (`kpi_report`): the navy slide, the panel
cards, the cyan rule under the header, the Huawei wordmark, the cover.

A chart is drawn from the very lines the app drew — the same series, the same
Draw Data colours, the same hours — as a PowerPoint chart the user can edit;
where the app hands over a picture of its own chart (`png`), that picture is
placed instead.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime

import pandas as pd

from rfopt.reports.kpi_report import (CYAN, LINE_COLOURS, MUTED, PANEL, Assets, _bg,
                                      _label_skip, _no_data, _panel, _rgb, _style_chart, _text)

PER_SLIDE = 6                       # three across, two down
MAX_LEGEND = 14                     # more names than this and a legend eats the slide
GRID_PX = (1080, 624)               # a picture for a box of the grid, its shape
WHOLE_PX = (1800, 738)              # and one for the combined chart's slide
ACROSS, DOWN = 3, 2
SLIDE_W, SLIDE_H = 13.333, 7.5
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@dataclass
class Deck:
    """What the deck says about itself, on the cover and in every header."""
    kpi: str
    tech: str
    by: str                          # "Per Site" / "Per Sector"
    window: str
    charts: int
    project: str = "R5 · Asiacell"
    prepared: str = ""

    @property
    def title(self) -> str:
        return f"{self.tech} {self.kpi}"

    def rows(self) -> list:
        return [("KPI", self.kpi), ("Technology", self.tech), ("Drawn by", self.by),
                ("Charts", f"{self.charts}"), ("Window", self.window),
                ("Project", self.project),
                ("Prepared", self.prepared or f"{datetime.now():%d %b %Y %H:%M}")]


def build(charts: list, combined, deck: Deck, assets: Assets) -> bytes:
    """`charts`: (title, panel, png or None); `combined`: the same, or None."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    sheets = (len(charts) + PER_SLIDE - 1) // PER_SLIDE
    pages = 1 + sheets + (1 if combined else 0)
    _cover(prs.slides.add_slide(blank), deck, assets)
    page = 2
    for start in range(0, len(charts), PER_SLIDE):
        _grid(prs.slides.add_slide(blank), charts[start:start + PER_SLIDE], deck, assets,
              page, pages, first=start + 1)
        page += 1
    if combined:
        _whole(prs.slides.add_slide(blank), combined, deck, assets, page, pages)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# the slides
# --------------------------------------------------------------------------- #
def _cover(slide, deck: Deck, assets: Assets) -> None:
    from pptx.util import Inches
    _bg(slide)
    if assets.cover and assets.cover.is_file():
        slide.shapes.add_picture(str(assets.cover), 0, 0, width=Inches(SLIDE_W),
                                 height=Inches(SLIDE_H))
    if assets.logo and assets.logo.is_file():
        slide.shapes.add_picture(str(assets.logo), Inches(0.55), Inches(0.45),
                                 height=Inches(0.55))
    _text(slide, 0.55, 1.2, 8.4, 1.5, [deck.title, "Analysis Report"], size=36, bold=True,
          colour="FFFFFF")
    y = 2.95
    for k, (label, value) in enumerate(deck.rows()):
        _text(slide, 0.58, y + k * 0.3, 1.9, 0.3, label, size=12.5, colour="CBD5E1")
        _text(slide, 2.45, y + k * 0.3, 0.25, 0.3, ":", size=12.5, colour="CBD5E1")
        _text(slide, 2.65, y + k * 0.3, 5.6, 0.3, value, size=12.5, bold=True, colour="FFFFFF")


def _head(slide, deck: Deck, assets: Assets, title: str, page: int, pages: int) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    _bg(slide)
    _text(slide, 0.45, 0.28, 9.4, 0.55, title, size=24, bold=True, colour="F8FAFC")
    _text(slide, 0.45, 0.8, 9.6, 0.35, f"{deck.kpi} · {deck.tech} · {deck.by} · {deck.window}",
          size=11.5, colour=MUTED)
    if assets.logo and assets.logo.is_file():
        slide.shapes.add_picture(str(assets.logo), Inches(10.9), Inches(0.33),
                                 height=Inches(0.46))
    bar = slide.shapes.add_connector(1, Inches(0.45), Inches(1.22), Inches(12.88), Inches(1.22))
    bar.line.color.rgb = _rgb(CYAN)
    bar.line.width = Pt(1.25)
    _text(slide, 11.9, 7.02, 1.0, 0.3, f"{page} / {pages}", size=10, colour=MUTED,
          align=PP_ALIGN.RIGHT)


def _grid(slide, part: list, deck: Deck, assets: Assets, page: int, pages: int,
          first: int) -> None:
    """Six charts on one slide: three across, two down, each under its title."""
    last = first + len(part) - 1
    _head(slide, deck, assets, f"Charts {first}–{last}" if last > first else f"Chart {first}",
          page, pages)
    x0, y0, gap = 0.45, 1.42, 0.16
    w = (12.43 - gap * (ACROSS - 1)) / ACROSS
    h = (5.72 - gap * (DOWN - 1)) / DOWN
    for k, (title, panel, png) in enumerate(part):
        row, col = divmod(k, ACROSS)
        x, y = x0 + col * (w + gap), y0 + row * (h + gap)
        _panel(slide, x, y, w, h)
        _text(slide, x + 0.14, y + 0.07, w - 0.28, 0.32, title, size=13, bold=True,
              colour="F1F5F9")
        _chart(slide, x + 0.1, y + 0.45, w - 0.2, h - 0.56, panel, png, legend=False, size=7)


def _whole(slide, one, deck: Deck, assets: Assets, page: int, pages: int) -> None:
    """The combined chart, a slide to itself so its cells stay readable."""
    title, panel, png = one
    lines = len(getattr(panel, "lines", {}) or {})
    named = 0 < lines <= MAX_LEGEND
    _head(slide, deck, assets, "Combined Chart", page, pages)
    _panel(slide, 0.45, 1.42, 12.43, 5.72)
    _text(slide, 0.6, 1.5, 12.1, 0.35, title, size=14, bold=True, colour="F1F5F9")
    if not named:
        # 285 names would leave no room for the lines themselves
        _text(slide, 0.6, 1.86, 12.1, 0.28, f"{lines:,} cells", size=10, colour=MUTED)
    top = 1.95 if named else 2.16
    _chart(slide, 0.6, top, 12.13, 7.14 - top, panel, png, legend=named, size=9)


# --------------------------------------------------------------------------- #
# a chart
# --------------------------------------------------------------------------- #
def _chart(slide, x, y, w, h, panel, png, *, legend: bool, size: int) -> None:
    if png:
        _picture(slide, x, y, w, h, png)
        return
    if panel is None or not getattr(panel, "lines", None):
        _no_data(slide, x, y, w, h, "No data for this chart.")
        return
    _native(slide, x, y, w, h, panel, legend=legend, size=size)


def _picture(slide, x, y, w, h, png: bytes) -> None:
    """The app's own picture of the chart, whole inside the box."""
    from pptx.util import Emu, Inches
    pic = slide.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y))
    scale = min(Inches(w) / pic.width, Inches(h) / pic.height)
    pic.width, pic.height = Emu(int(pic.width * scale)), Emu(int(pic.height * scale))
    pic.left = Inches(x) + Emu(int((Inches(w) - pic.width) / 2))
    pic.top = Inches(y) + Emu(int((Inches(h) - pic.height) / 2))


def _native(slide, x, y, w, h, panel, *, legend: bool, size: int) -> None:
    """The panel's own lines as a PowerPoint chart, in Draw Data's colours."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
    from pptx.util import Inches, Pt

    hours = None
    for s in panel.lines.values():
        hours = s.index if hours is None else hours.union(s.index)
    data = CategoryChartData()
    data.categories = [f"{t:%d %b %H:%M}" for t in hours]
    for name, s in panel.lines.items():
        vals = s.reindex(hours).to_numpy(dtype=float)
        data.add_series(str(name), [None if pd.isna(v) else float(v) for v in vals])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w),
                                Inches(h), data)
    chart = gf.chart
    _style_chart(chart, legend=legend, size=size)
    for i, series in enumerate(chart.plots[0].series):
        series.format.line.color.rgb = _rgb(LINE_COLOURS[i % len(LINE_COLOURS)].lstrip("#"))
        series.format.line.width = Pt(1.5 if legend else 1.0)
        series.smooth = False
        series.marker.style = XL_MARKER_STYLE.NONE
    step = max(1, len(hours) // (10 if legend else 4))
    if step > 1:
        _label_skip(chart.category_axis._element, step)
    _dark(chart)


def _dark(chart) -> None:
    """The slide's navy under the chart, so it sits in its card like the app's."""
    space = chart._chartSpace
    plot = space.find(f"{_C}chart/{_C}plotArea")
    _fill(space, PANEL, after=space.find(f"{_C}chart"))
    if plot is not None:
        _fill(plot, None)                      # the card shows through the plot


def _fill(parent, colour: str | None, after=None) -> None:
    """A c:spPr of its own for `parent`: `colour`, or nothing at all."""
    sp = parent.makeelement(f"{_C}spPr", {})
    if colour:
        solid = parent.makeelement(f"{_A}solidFill", {})
        rgb = parent.makeelement(f"{_A}srgbClr", {"val": colour})
        solid.append(rgb)
        sp.append(solid)
    else:
        sp.append(parent.makeelement(f"{_A}noFill", {}))
    line = parent.makeelement(f"{_A}ln", {})
    line.append(parent.makeelement(f"{_A}noFill", {}))
    sp.append(line)
    old = parent.find(f"{_C}spPr")
    if old is not None:
        parent.remove(old)
    if after is not None:
        after.addnext(sp)
    else:
        parent.append(sp)


__all__ = ["Deck", "PER_SLIDE", "build"]
