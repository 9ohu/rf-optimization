"""Annotated Excel workbook: overview, diagnoses, recommendations, KPIs."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

_SEV_FILL = {"Critical": "FFC7CE", "critical": "FFC7CE",
             "Warning": "FFEB9C", "warning": "FFEB9C",
             "P1": "FFC7CE", "P2": "FFD9A5", "P3": "FFF2CC", "P4": "E2EFDA"}


def _autosize(ws, df: pd.DataFrame, cap: int = 60) -> None:
    from openpyxl.utils import get_column_letter
    for i, col in enumerate(df.columns, start=1):
        width = max(len(str(col)),
                    int(df[col].astype(str).str.len().quantile(0.9) or 10)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, cap)


def _style_header(ws) -> None:
    from openpyxl.styles import Font, PatternFill
    fill = PatternFill("solid", fgColor="305496")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    ws.freeze_panes = "A2"


def _shade(ws, df: pd.DataFrame, col_name: str) -> None:
    from openpyxl.styles import PatternFill
    if col_name not in df.columns:
        return
    j = list(df.columns).index(col_name) + 1
    for i, val in enumerate(df[col_name].astype(str), start=2):
        hexc = _SEV_FILL.get(val)
        if hexc:
            ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=hexc)


def write_excel_report(result, out_path: str | Path,
                       narratives: list | None = None) -> Path:
    """Write a multi-sheet workbook from an AnalysisResult."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    s = result.summary()
    overview = pd.DataFrame(
        [{"metric": k, "value": v} for k, v in s.items()]
        + [{"metric": "generated", "value": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M")}]
    )

    diag_df = result.diagnoses_df()
    rec_df = result.recommendations_df()
    br_df = result.breaches_df()
    tr_df = result.trends_df()

    by_class = (diag_df.groupby(["problem_class", "severity"]).size()
                .unstack(fill_value=0).reset_index()
                if not diag_df.empty else pd.DataFrame())
    by_prio = (rec_df.groupby("priority").size().reset_index(name="count")
               if not rec_df.empty else pd.DataFrame())

    # full recommendation text for the action plan
    plan_rows = []
    for r in result.recommendations:
        plan_rows.append({
            "priority": r.priority, "cell": r.cell_id, "site": r.site_id,
            "problem": r.problem, "category": r.category,
            "confidence": round(r.confidence, 2),
            "action": r.action_type, "parameter": r.parameter,
            "current": r.current_value, "recommended": r.recommended_value,
            "root_cause": r.root_cause,
            "evidence": " | ".join(r.evidence),
            "expected_impact": r.expected_impact,
            "risks": " | ".join(r.risks),
            "monitor": ", ".join(r.monitor_kpis),
            "calculation": " ; ".join(r.math_notes),
            "alternatives": " | ".join(r.alternatives),
        })
    plan_df = pd.DataFrame(plan_rows)

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        overview.to_excel(xw, sheet_name="Overview", index=False)
        if not by_class.empty:
            by_class.to_excel(xw, sheet_name="Overview", index=False,
                              startrow=len(overview) + 3)
        if not by_prio.empty:
            by_prio.to_excel(xw, sheet_name="Overview", index=False,
                             startrow=len(overview) + 3, startcol=8)
        (plan_df if not plan_df.empty else pd.DataFrame({"info": ["no actions"]})
         ).to_excel(xw, sheet_name="Action Plan", index=False)
        (diag_df if not diag_df.empty else pd.DataFrame({"info": ["none"]})
         ).to_excel(xw, sheet_name="Diagnoses", index=False)
        (br_df if not br_df.empty else pd.DataFrame({"info": ["none"]})
         ).to_excel(xw, sheet_name="KPI Breaches", index=False)
        (tr_df if not tr_df.empty else pd.DataFrame({"info": ["none"]})
         ).to_excel(xw, sheet_name="Trends", index=False)
        result.agg_cell.to_excel(xw, sheet_name="Cell KPIs", index=False)
        if not result.agg_site.empty:
            result.agg_site.to_excel(xw, sheet_name="Site KPIs", index=False)
        if not result.busy_hour.empty:
            result.busy_hour.to_excel(xw, sheet_name="Busy Hour", index=False)
        if narratives:
            nrows = [{"cell": n.entity_id, "problem": n.problem,
                      "priority": n.priority, "source": n.source,
                      **{k: v for k, v in n.sections.items()}}
                     for n in narratives]
            pd.DataFrame(nrows).to_excel(xw, sheet_name="AI Narratives",
                                         index=False)

    # post-format
    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    fmt = {
        "Action Plan": (plan_df, "priority"),
        "Diagnoses": (diag_df, "severity"),
        "KPI Breaches": (br_df, "severity"),
    }
    for name in wb.sheetnames:
        ws = wb[name]
        _style_header(ws)
        if name in fmt and not fmt[name][0].empty:
            df, col = fmt[name]
            _autosize(ws, df)
            _shade(ws, df, col)
    wb.save(out_path)
    return out_path
