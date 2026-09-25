"""The KPI report: one judged KPI over an area and a period, and its exports.

`KpiReport` holds what the KPI Analysis judgement found (built by the app's
`_kpi_report` from `_kpi_health.object_values` with the user's threshold, and
`rfopt.kpi.anomaly.sudden_spikes` against each cell's own normal behaviour, so
no KPI rule lives here) and derives everything the report shows from it: the
totals, the affected cells per governorate / Sup District, the worst sites and
their hourly KPI, the hourly network trend and the tickets. The preview, the
PowerPoint (`build_pptx`) and the Excel workbook (`build_xlsx`) all read the
same object, so they cannot disagree.

* A cell is affected when it is beyond the threshold (in the KPI's own bad
  direction) or shows a sudden spike; the charts count only affected cells.
* The Excel workbook is the Draw Data export itself
  (`kpi_pivot.write_pivot_workbook`: a row per cell, a column per hour), of
  the affected cells.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

NAVY, PANEL, LINE = "071525", "0B1F33", "1E3A5F"
TEXT, MUTED, CYAN, BLUE = "E2E8F0", "94A3B8", "20BFFF", "1597FF"
RED, GREEN, AMBER = "EF4444", "22C55E", "FACC15"
TICKET_COLS = ["Ticket ID", "Site ID", "City", "SLA Target Time", "Create Time",
               "Problem Time", "Is CMC", "KPI Value"]
SLIDES = ("Cover", "R5 Network Trend", "Affected Cells by Area", "Top Sites",
          "Tickets", "Received Tickets")
# the Draw Data chart's line colours (app/_charts.LINE_COLOURS), in its order
LINE_COLOURS = ["#20BFFF", "#A78BFA", "#FB923C", "#F472B6", "#4ADE80",
                "#FACC15", "#2DD4BF", "#E879F9", "#CBD5E1", "#F87171"]


@dataclass
class KpiReport:
    tech: str                              # "4G" / "3G"
    kpi: str                               # "4G TDD Interference"
    column: str                            # the export's own column of the KPI
    unit: str                              # " dBm", "%", ""
    low_is_bad: bool                       # True: a value below the threshold is the problem
    per_day: bool                          # the threshold is on a 24 h total
    threshold: float
    governorate: str                       # "" = every governorate
    sup_districts: list
    site: str                              # "" = every site
    start: date
    end: date
    # one row per cell (4G) / NodeB (3G) in scope: object_id, site_id, cell_name,
    # governorate, sup_district, value, peak, peak_time, worst (its worst hour),
    # beyond (the threshold), spike, spike_value, spike_time, affected
    cells: pd.DataFrame
    rows: pd.DataFrame                     # their hourly values, as the KPI export has them
    trend: pd.Series                       # hour -> network average
    tickets: pd.DataFrame                  # TICKET_COLS, in scope
    generated: datetime = field(default_factory=datetime.now)

    # ---------------------------------------------------------------- scope
    @property
    def gov_text(self) -> str:
        return self.governorate or "All governorates (R5)"

    @property
    def sd_text(self) -> str:
        return ", ".join(self.sup_districts) if self.sup_districts else "All Sup Districts"

    @property
    def site_text(self) -> str:
        return self.site or "All Sites"

    @property
    def period_text(self) -> str:
        return (f"{self.start:%d %b %Y}" if self.start == self.end
                else f"{self.start:%d %b %Y} – {self.end:%d %b %Y}")

    @property
    def op(self) -> str:
        return "<" if self.low_is_bad else ">"

    @property
    def threshold_text(self) -> str:
        return (f"{self.op} {self.threshold:g}{self.unit}"
                + (" per 24 h" if self.per_day else ""))

    # --------------------------------------------------------------- totals
    @property
    def affected(self) -> pd.DataFrame:
        """The cells beyond the threshold or with a sudden spike."""
        return self.cells[self.cells["affected"]]

    @property
    def total_sites(self) -> int:
        return int(self.cells["site_id"].nunique())

    @property
    def affected_sites(self) -> int:
        return int(self.affected["site_id"].nunique())

    @property
    def total_cells(self) -> int:
        return int(len(self.cells))

    @property
    def affected_cells(self) -> int:
        return int(len(self.affected))

    @property
    def spike_cells(self) -> int:
        return int(self.cells["spike"].sum()) if len(self.cells) else 0

    @property
    def total_tickets(self) -> int:
        return int(len(self.tickets))

    @staticmethod
    def pct(n: int, of: int) -> str:
        return f"{100.0 * n / of:.1f}%" if of else "0.0%"

    # ------------------------------------------- charts: affected records only
    def affected_by(self, column: str, keep=()) -> pd.Series:
        """Affected cells per area, worst first; `keep` areas (the Sup Districts
        picked) stay listed with 0."""
        a = self.affected
        s = a.groupby(column).size() if len(a) else pd.Series(dtype=int)
        s = s.drop(["", "Not in EP tracker"], errors="ignore")
        for k in keep:
            if k not in s.index:
                s.loc[k] = 0
        return s.astype(int).sort_values(ascending=False, kind="stable")

    def top_sites(self, n: int = 10) -> pd.DataFrame:
        """The affected sites, each at its worst cell, worst first in the KPI's
        own bad direction: by the worst hour, then by the value over the period."""
        a = self.affected
        if a.empty:
            return pd.DataFrame(columns=["site_id", "worst", "value", "cell_name"])
        sign = 1.0 if self.low_is_bad else -1.0            # ascending = worst first
        a = a.assign(_w=sign * a["worst"].astype(float), _v=sign * a["value"].astype(float))
        a = a.sort_values(["_w", "_v"], kind="stable").drop_duplicates("site_id")
        return a[["site_id", "worst", "value", "cell_name"]].head(n).reset_index(drop=True)

    def site_lines(self, n: int = 10) -> dict:
        """The top sites' KPI hour by hour, one line per site, drawn the way Draw
        Data draws a site: its cells averaged (a level) or added up (a count)."""
        from rfopt.kpi.trends import series_for
        out = {}
        for site in self.top_sites(n)["site_id"]:
            line = series_for(self.rows, self.column, level="Site", obj=site)
            if len(line):
                out[str(site)] = line
        return out

    def top_ticket_sites(self, n: int = 10) -> pd.Series:
        t = self.tickets
        if t.empty:
            return pd.Series(dtype=int)
        s = t.loc[t["Site ID"].astype(str).str.strip() != "", "Site ID"].value_counts()
        return s.head(n)

    # ------------------------------------------------------------ narrative
    @property
    def description(self) -> str:
        """A short technical statement of what the data shows — nothing else."""
        where = (", ".join(self.sup_districts) if self.sup_districts
                 else self.governorate or "R5")
        if not self.total_cells:
            return (f"No {self.kpi} data for {where} in {self.period_text}: the loaded KPI "
                    "exports do not cover this area and period.")
        cells = f"{self.affected_cells:,} of {self.total_cells:,} cells"
        if self.tech == "3G":
            cells = f"{self.affected_cells:,} of {self.total_cells:,} NodeBs"
        jump = "drop" if self.low_is_bad else "spike"
        if not self.affected_cells:
            return (f"{self.kpi} stayed within the configured threshold, with no sudden "
                    f"{jump}, on all {self.total_cells:,} "
                    f"{'NodeBs' if self.tech == '3G' else 'cells'} of {where} during "
                    f"{self.period_text}.")
        verb = "fell below" if self.low_is_bad else "exceeded"
        what = f"{verb} the configured threshold" + (
            f" or showed a sudden {jump}" if self.spike_cells else "")
        top = self.affected_by("sup_district").head(2)
        mainly = (f", mainly in {' and '.join(top.index)}" if len(top) and not self.site
                  else f" at {self.site}" if self.site else "")
        return (f"{self.kpi} {what} on {cells} "
                f"({self.pct(self.affected_cells, self.total_cells)}) at "
                f"{self.affected_sites:,} of {self.total_sites:,} sites during "
                f"{self.period_text}{mainly}.")

    def summary_rows(self) -> list[tuple[str, str]]:
        """The cover's summary: no threshold here (it stays in the charts and Excel)."""
        return [("Governorate", self.gov_text), ("Sup District(s)", self.sd_text),
                ("Site", self.site_text), ("Technology", self.tech), ("KPI", self.kpi),
                ("Period", self.period_text), ("Total Sites", f"{self.total_sites:,}"),
                ("Affected Sites", f"{self.affected_sites:,} "
                                   f"({self.pct(self.affected_sites, self.total_sites)})"),
                ("Total Cells" if self.tech != "3G" else "Total NodeBs", f"{self.total_cells:,}"),
                ("Affected Cells" if self.tech != "3G" else "Affected NodeBs",
                 f"{self.affected_cells:,} ({self.pct(self.affected_cells, self.total_cells)})"),
                ("Total Tickets", f"{self.total_tickets:,}")]

    # --------------------------------------------------------------- Excel
    def excel_rows(self) -> pd.DataFrame:
        """What the workbook holds: the affected cells' hourly values."""
        keep = set(self.affected["object_id"].astype(str))
        return self.rows[self.rows["object"].astype(str).isin(keep)]

    @property
    def sheet(self) -> str:
        """The workbook's sheet: the Draw Data export's name for the KPI."""
        from rfopt.reports.kpi_pivot import spec_for
        return spec_for(self.column).sheet

    def excel_preview(self) -> pd.DataFrame:
        """The sheet as the workbook shows it: Row Labels, the duplex / RNC, and
        one column per hour (`kpi_pivot.pivot_one`)."""
        from rfopt.reports.kpi_pivot import pivot_one
        piv = pivot_one(self.excel_rows(), self.column)
        if piv.empty:
            return piv
        n_id = sum(1 for c in piv.columns if not isinstance(c, pd.Timestamp))
        piv.columns = (["Row Labels"] + [str(c) for c in piv.columns[1:n_id]]
                       + [pd.Timestamp(h).strftime("%m/%d %I:%M %p")
                          for h in piv.columns[n_id:]])
        return piv


# --------------------------------------------------------------------------- #
# file names
# --------------------------------------------------------------------------- #
def _slug(text: str) -> str:
    s = re.sub(r'[\\/:*?"<>|\x00-\x1f]', " ", str(text))
    s = re.sub(r"[^\w\-\.]+", "_", s.strip(), flags=re.UNICODE)
    return re.sub(r"_+", "_", s).strip("_.")


def file_stem(r: KpiReport, today: date | None = None) -> str:
    """KPI_Report_4G_TDD_Interference_Basrah_Zubair_2026-09-18."""
    kpi = r.kpi[len(r.tech):].strip() if r.kpi.startswith(r.tech) else r.kpi
    parts = ["KPI_Report", r.tech, kpi, r.governorate or "R5"]
    if r.sup_districts:
        parts += list(r.sup_districts[:2])
        if len(r.sup_districts) > 2:
            parts.append(f"plus{len(r.sup_districts) - 2}")
    if r.site:
        parts.append(r.site)
    parts.append(f"{today or date.today():%Y-%m-%d}")
    return _slug("_".join(_slug(p) for p in parts if p))[:150]


def pptx_name(r: KpiReport, today: date | None = None) -> str:
    return file_stem(r, today) + ".pptx"


def xlsx_name(r: KpiReport, today: date | None = None) -> str:
    return file_stem(r, today).replace("KPI_Report_", "KPI_Report_Data_", 1) + ".xlsx"


# --------------------------------------------------------------------------- #
# Excel: the Draw Data export, of the affected cells
# --------------------------------------------------------------------------- #
def build_xlsx(r: KpiReport) -> bytes:
    """The workbook Draw Data's Excel export writes (`write_pivot_workbook`,
    unchanged: one sheet per KPI, a row per cell, the hours across), holding
    the report's affected cells over its period."""
    import tempfile
    from rfopt.reports.kpi_pivot import write_pivot_workbook
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "report.xlsx"
        write_pivot_workbook(r.excel_rows(), [r.column], path)
        return path.read_bytes()


# --------------------------------------------------------------------------- #
# PowerPoint: real slides, native (editable) charts
# --------------------------------------------------------------------------- #
@dataclass
class Assets:
    cover: Path          # the Basrah bridge, veiled on the left for the text
    logo: Path           # the Huawei logo, white wordmark (dark slides)


def build_pptx(r: KpiReport, assets: Assets) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    _cover(prs.slides.add_slide(blank), r, assets)
    _trend_slide(prs.slides.add_slide(blank), r, assets, 2)
    _areas_slide(prs.slides.add_slide(blank), r, assets, 3)
    _top_sites_slide(prs.slides.add_slide(blank), r, assets, 4)
    _tickets_slide(prs.slides.add_slide(blank), r, assets, 5)
    _ticket_table_slide(prs.slides.add_slide(blank), r, assets, 6)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _rgb(hex_: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_)


def _bg(slide, colour: str = NAVY) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(colour)


def _text(slide, x, y, w, h, text, size=14, bold=False, colour=TEXT, align=None,
          anchor=None, font="Segoe UI"):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.vertical_anchor = anchor or MSO_ANCHOR.TOP
    lines = text if isinstance(text, list) else [text]
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align or PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, font
        run.font.color.rgb = _rgb(colour)
    return box


def _panel(slide, x, y, w, h, colour: str = PANEL):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    shp.adjustments[0] = 0.04
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(colour)
    shp.line.color.rgb = _rgb(LINE)
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def _header(slide, r: KpiReport, assets: Assets, title: str, page: int) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    _bg(slide)
    _text(slide, 0.45, 0.28, 9.4, 0.55, title, size=24, bold=True, colour="F8FAFC")
    _text(slide, 0.45, 0.8, 9.6, 0.35,
          f"{r.kpi} · {r.gov_text} · {r.sd_text if r.sup_districts else r.site_text} · "
          f"{r.period_text}", size=11.5, colour=MUTED)
    slide.shapes.add_picture(str(assets.logo), Inches(10.9), Inches(0.33), height=Inches(0.46))
    bar = slide.shapes.add_connector(1, Inches(0.45), Inches(1.22), Inches(12.88), Inches(1.22))
    bar.line.color.rgb = _rgb(CYAN)
    bar.line.width = Pt(1.25)
    _text(slide, 11.9, 7.02, 1.0, 0.3, f"{page} / {len(SLIDES)}", size=10, colour=MUTED,
          align=PP_ALIGN.RIGHT)


def _cover(slide, r: KpiReport, assets: Assets) -> None:
    from pptx.util import Inches
    slide.shapes.add_picture(str(assets.cover), 0, 0, width=Inches(13.333),
                             height=Inches(7.5))
    slide.shapes.add_picture(str(assets.logo), Inches(0.55), Inches(0.45), height=Inches(0.55))
    title = r.kpi[len(r.tech):].strip() if r.kpi.startswith(r.tech) else r.kpi
    _text(slide, 0.55, 1.2, 7.6, 1.5, [f"{r.tech} {title}", "Analysis Report"], size=36,
          bold=True, colour="FFFFFF")
    rows = r.summary_rows()
    y = 2.95
    for k, (label, value) in enumerate(rows):
        _text(slide, 0.58, y + k * 0.3, 1.9, 0.3, label, size=12.5, colour="CBD5E1")
        _text(slide, 2.45, y + k * 0.3, 0.25, 0.3, ":", size=12.5, colour="CBD5E1")
        _text(slide, 2.65, y + k * 0.3, 5.2, 0.3, value, size=12.5, bold=True, colour="FFFFFF")
    _text(slide, 0.58, y + len(rows) * 0.3 + 0.18, 6.6, 0.8, r.description, size=12,
          colour="E2E8F0")


def _style_chart(chart, legend: bool = False, size: int = 10) -> None:
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Pt
    chart.font.size = Pt(size)
    chart.font.color.rgb = _rgb("CBD5E1")
    chart.font.name = "Segoe UI"
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.color.rgb = _rgb("CBD5E1")
    for axis in (chart.category_axis, chart.value_axis):
        axis.format.line.color.rgb = _rgb(LINE)
        axis.tick_labels.font.color.rgb = _rgb("CBD5E1")
        axis.tick_labels.font.size = Pt(size)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb("16324F")


def _chart_title(slide, x, y, w, text, note: str = "") -> None:
    _text(slide, x + 0.15, y + 0.08, w - 0.3, 0.35, text, size=13.5, bold=True, colour="F1F5F9")
    if note:
        _text(slide, x + 0.15, y + 0.4, w - 0.3, 0.28, note, size=10, colour=MUTED)


def _no_data(slide, x, y, w, h, text: str) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    _text(slide, x, y, w, h, text, size=12, colour=MUTED, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)


def _bars(slide, x, y, w, h, series: pd.Series, *, horizontal: bool = False,
          colour: str = BLUE, first: str | None = RED, fmt: str = "0",
          name: str = "Affected cells") -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Inches, Pt
    data = CategoryChartData()
    data.categories = [str(c) for c in series.index]
    data.add_series(name, [float(v) for v in series.to_numpy()])
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = gf.chart
    _style_chart(chart)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(9.5)
    plot.data_labels.font.color.rgb = _rgb("F1F5F9")
    plot.data_labels.number_format = fmt
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = _rgb(colour)
    if first and len(series):
        pt = ser.points[0]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(first)
    if horizontal:
        chart.category_axis.reverse_order = True      # the first (worst) at the top
    chart.value_axis.tick_labels.number_format = fmt
    chart.value_axis.tick_labels.number_format_is_linked = False


def trend_note(r: KpiReport) -> str:
    """Under the hourly trend: the threshold, or why a 24 h one is not drawn."""
    return f"Threshold {r.threshold_text}" + (
        " · a 24 h total, not drawn on the hourly chart" if r.per_day else "")


def _trend_slide(slide, r: KpiReport, assets: Assets, page: int) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.util import Inches, Pt
    _header(slide, r, assets, "R5 Network Trend", page)
    _panel(slide, 0.45, 1.45, 12.43, 5.35)
    _chart_title(slide, 0.45, 1.45, 12.43, f"{r.kpi} — network average (hourly)",
                 trend_note(r))
    t = r.trend.dropna()
    if t.empty:
        _no_data(slide, 0.45, 2.2, 12.43, 4.4, "No KPI data in the selected area and period.")
    else:
        data = CategoryChartData()
        data.categories = [f"{x:%d %b %H:%M}" for x in t.index]
        data.add_series("Network Average", [float(v) for v in t.to_numpy()])
        if not r.per_day:
            data.add_series(f"Threshold ({r.threshold:g}{r.unit})",
                            [float(r.threshold)] * len(t))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.65), Inches(2.25),
                                    Inches(12.0), Inches(4.4), data)
        chart = gf.chart
        _style_chart(chart, legend=True)
        avg = chart.plots[0].series[0]
        avg.format.line.color.rgb = _rgb(CYAN)
        avg.format.line.width = Pt(2.25)
        avg.smooth = False
        avg.marker.style = XL_MARKER_STYLE.CIRCLE if len(t) <= 60 else XL_MARKER_STYLE.NONE
        avg.marker.size = 5
        avg.marker.format.fill.solid()
        avg.marker.format.fill.fore_color.rgb = _rgb(CYAN)
        if not r.per_day:
            thr = chart.plots[0].series[1]
            thr.format.line.color.rgb = _rgb(RED)
            thr.format.line.width = Pt(1.75)
            thr.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            thr.marker.style = XL_MARKER_STYLE.NONE
            thr.smooth = False
        step = max(1, len(t) // 12)
        if step > 1:
            _label_skip(chart.category_axis._element, step)
    _text(slide, 0.6, 6.9, 11.2, 0.3, r.description, size=10.5, colour=MUTED)


def _label_skip(ax, step: int) -> None:
    """Every `step`-th category label only: c:tickLblSkip, in its schema place
    (after c:lblOffset, before c:tickMarkSkip / c:noMultiLvlLbl)."""
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    skip = ax.makeelement(ns + "tickLblSkip", {"val": str(step)})
    after = [ns + n for n in ("tickMarkSkip", "noMultiLvlLbl", "extLst")]
    nxt = next((c for c in ax if c.tag in after), None)
    if nxt is not None:
        nxt.addprevious(skip)
    else:
        ax.append(skip)


def _areas_slide(slide, r: KpiReport, assets: Assets, page: int) -> None:
    _header(slide, r, assets, "Affected Cells by Area", page)
    what = "NodeBs" if r.tech == "3G" else "cells"
    panels = (("Governorate", "governorate", ()),
              ("Sup District", "sup_district", tuple(r.sup_districts)))
    w = 6.13
    for k, (title, col, keep) in enumerate(panels):
        x = 0.45 + k * (w + 0.17)
        _panel(slide, x, 1.45, w, 5.35)
        s = r.affected_by(col, keep).head(8)
        _chart_title(slide, x, 1.45, w, f"Affected {what} by {title}",
                     f"beyond the threshold {r.threshold_text} or with a sudden spike")
        if s.empty or not s.sum():
            _no_data(slide, x, 2.3, w, 4.2, f"No {what} beyond the threshold.")
            continue
        _bars(slide, x + 0.1, 2.2, w - 0.2, 4.5, s, name=f"Affected {what}")


def top_sites_note(r: KpiReport) -> str:
    """Under the Top Sites chart: which sites, and how a site's hour is drawn."""
    return ("the 10 worst affected sites (beyond " + r.threshold_text + " or a sudden spike) "
            "· hourly · a site as Draw Data draws it: its cells "
            + ("added up" if _agg(r.column) == "sum" else "averaged"))


def _agg(column: str) -> str:
    from rfopt.kpi.trends import agg_how
    return agg_how(column)


def _top_sites_slide(slide, r: KpiReport, assets: Assets, page: int) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.util import Inches, Pt
    _header(slide, r, assets, "Top Sites", page)
    kpi_short = r.kpi[len(r.tech):].strip() if r.kpi.startswith(r.tech) else r.kpi
    _panel(slide, 0.45, 1.45, 12.43, 5.35)
    _chart_title(slide, 0.45, 1.45, 12.43, f"Top 10 Sites by {kpi_short}", top_sites_note(r))
    lines = r.site_lines(10)
    if not lines:
        _no_data(slide, 0.45, 2.3, 12.43, 4.2, "No affected site in the area and period.")
        return
    hours = sorted(set().union(*(set(s.index) for s in lines.values())))
    data = CategoryChartData()
    data.categories = [f"{t:%d %b %H:%M}" for t in hours]
    for site, s in lines.items():
        data.add_series(site, [None if pd.isna(v) else float(v)
                               for v in s.reindex(hours).to_numpy()])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(2.2), Inches(12.1),
                                Inches(4.55), data)
    chart = gf.chart
    _style_chart(chart, legend=True)
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    for k, ser in enumerate(chart.plots[0].series):
        ser.format.line.color.rgb = _rgb(LINE_COLOURS[k % len(LINE_COLOURS)].lstrip("#"))
        ser.format.line.width = Pt(1.6)
        ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
    step = max(1, len(hours) // 12)
    if step > 1:
        _label_skip(chart.category_axis._element, step)


def _tickets_slide(slide, r: KpiReport, assets: Assets, page: int) -> None:
    from pptx.enum.text import PP_ALIGN
    _header(slide, r, assets, "Tickets", page)
    _panel(slide, 0.45, 1.45, 3.4, 5.35)
    _text(slide, 0.6, 1.6, 3.1, 0.4, "Total Tickets", size=14, bold=True, colour="F1F5F9")
    _text(slide, 0.6, 2.2, 3.1, 1.2, f"{r.total_tickets:,}", size=60, bold=True, colour=CYAN,
          align=PP_ALIGN.LEFT)
    _text(slide, 0.6, 3.6, 3.1, 2.8,
          [f"Area: {r.gov_text}", f"Sup District(s): {r.sd_text}", f"Site: {r.site_text}",
           f"Problem time within {r.period_text}", "Source: the Daily Target (Complaint Data)"],
          size=11, colour=MUTED)
    _panel(slide, 4.0, 1.45, 8.88, 5.35)
    _chart_title(slide, 4.0, 1.45, 8.88, "Top 10 Sites by Number of Tickets",
                 "tickets per site in the selected area and period")
    s = r.top_ticket_sites(10)
    if s.empty:
        _no_data(slide, 4.0, 2.3, 8.88, 4.2, "No ticket in the selected area and period.")
    else:
        _bars(slide, 4.1, 2.2, 8.68, 4.5, s, horizontal=True, first=None,
              name="Tickets")


def ticket_header(r: KpiReport, col: str) -> str:
    """A ticket table heading: the KPI value carries the KPI's unit."""
    unit = r.unit.strip()
    return f"{col} ({unit})" if col == "KPI Value" and unit else col


def ticket_text(v) -> str:
    """A ticket table cell as the slide shows it."""
    if isinstance(v, (pd.Timestamp, datetime)):
        return "" if pd.isna(v) else f"{v:%Y-%m-%d %H:%M}"
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if isinstance(v, (float, np.floating)):
        return f"{v:,.0f}" if abs(v) >= 1000 or float(v).is_integer() else f"{v:,.2f}"
    return str(v)


def _ticket_table_slide(slide, r: KpiReport, assets: Assets, page: int) -> None:
    from pptx.util import Inches, Pt
    _header(slide, r, assets, "Received Tickets", page)
    t = r.tickets.sort_values("Problem Time", ascending=False, kind="stable").head(14)
    if t.empty:
        _panel(slide, 0.45, 1.45, 12.43, 5.35)
        _no_data(slide, 0.45, 2.3, 12.43, 4.2, "No ticket in the selected area and period.")
        return
    rows, cols = len(t) + 1, len(TICKET_COLS)
    shape = slide.shapes.add_table(rows, cols, Inches(0.45), Inches(1.5), Inches(12.43),
                                   Inches(0.36 * rows))
    table = shape.table
    widths = [2.35, 1.1, 1.55, 1.65, 1.65, 1.65, 0.85, 1.63]
    for j, wd in enumerate(widths):
        table.columns[j].width = Inches(wd)
    for j, c in enumerate(TICKET_COLS):
        cell = table.cell(0, j)
        cell.text = ticket_header(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(PANEL)
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size, para.runs[0].font.bold = Pt(11.5), True
        para.runs[0].font.color.rgb = _rgb(CYAN)
    for i, row in enumerate(t[TICKET_COLS].itertuples(index=False), start=1):
        for j, v in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ticket_text(v)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb("0D2945" if i % 2 else PANEL)
            para = cell.text_frame.paragraphs[0]
            if para.runs:
                para.runs[0].font.size = Pt(10)
                para.runs[0].font.color.rgb = _rgb(TEXT)
    more = r.total_tickets - len(t)
    if more > 0:
        _text(slide, 0.5, 1.55 + 0.36 * rows, 12, 0.3,
              f"… and {more:,} more ticket{'s' if more != 1 else ''} in the selected area "
              "and period.", size=10.5, colour=MUTED)
